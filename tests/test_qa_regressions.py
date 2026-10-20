"""Regressions for bugs found by the QA sweep.

1. GIF -> WebM/MP4 failed: a GIF decodes to an alpha pixel format (gbrap)
   that libvpx/x264 reject, so ffmpeg wrote nothing.
2. ODS/XLS -> TSV failed: LibreOffice's CSV filter always emits
   "<stem>.csv", so asking for .tsv never produced the expected file.
"""

from __future__ import annotations

import shutil
import subprocess

import pytest

HAVE_FFMPEG = shutil.which("ffmpeg") is not None
HAVE_SOFFICE = shutil.which("soffice") is not None


@pytest.fixture
def gif(tmp_path):
    if not HAVE_FFMPEG:
        pytest.skip("backend missing: ffmpeg")
    p = tmp_path / "clip.gif"
    subprocess.run(
        [
            "ffmpeg",
            "-y",
            "-v",
            "error",
            "-f",
            "lavfi",
            "-i",
            "testsrc=duration=1:size=64x48:rate=5",
            str(p),
        ],
        check=True,
        capture_output=True,
    )
    return p


@pytest.mark.parametrize("dest", ["webm", "mp4"])
def test_gif_converts_despite_alpha_pixel_format(gif, tmp_path, dest):
    from kumquat.converters_av import convert_media

    out = convert_media(gif, dest, tmp_path / "out")

    assert out.exists()
    assert out.stat().st_size > 0, f"{dest} output was empty"


@pytest.fixture
def ods(tmp_path):
    if not HAVE_SOFFICE:
        pytest.skip("backend missing: soffice")
    from kumquat._loshell import profile_args

    csv = tmp_path / "sheet.csv"
    csv.write_text("a,b\n1,2\n")
    subprocess.run(
        [
            "soffice",
            *profile_args(tmp_path / "isolated"),
            "--headless",
            "--convert-to",
            "ods",
            "--outdir",
            str(tmp_path),
            str(csv),
        ],
        check=True,
        capture_output=True,
        timeout=60,
    )
    p = tmp_path / "sheet.ods"
    assert p.is_file(), "installed soffice failed to generate a genuine ODS fixture"
    return p


def test_ods_to_tsv_writes_real_tsv(ods, tmp_path):
    from kumquat.converters_docs import convert_sheet

    out = convert_sheet(ods, "tsv", tmp_path / "out")

    assert out.suffix == ".tsv"
    assert out.exists() and out.stat().st_size > 0
    assert "\t" in out.read_text()


def test_ods_to_csv_still_works(ods, tmp_path):
    from kumquat.converters_docs import convert_sheet

    out = convert_sheet(ods, "csv", tmp_path / "out")

    assert out.suffix == ".csv"
    assert "," in out.read_text()


# --- epub in/out via pandoc (was hard-refused even with pandoc installed) ---

HAVE_PANDOC = shutil.which("pandoc") is not None


@pytest.fixture
def markdown(tmp_path):
    p = tmp_path / "note.md"
    p.write_text("# Title\n\nHello epub.\n")
    return p


def test_md_to_epub_uses_pandoc(markdown, tmp_path):
    if not HAVE_PANDOC:
        pytest.skip("backend missing: pandoc")
    from kumquat.converters_docs import convert_doc

    out = convert_doc(markdown, "epub", tmp_path / "out")

    assert out.suffix == ".epub"
    assert out.stat().st_size > 0
    assert out.read_bytes()[:2] == b"PK"  # epub is a zip container


def test_epub_input_converts_to_text(tmp_path):
    if not HAVE_PANDOC:
        pytest.skip("backend missing: pandoc")
    from kumquat.converters_docs import convert_doc

    md = tmp_path / "seed.md"
    md.write_text("# Title\n\nHello epub.\n")
    epub = convert_doc(md, "epub", tmp_path)
    out = convert_doc(epub, "txt", tmp_path / "back")

    assert out.suffix == ".txt"
    assert "Hello epub" in out.read_text()


# Full-sweep accounting and content gates must themselves detect bad outputs.
def test_full_sweep_inventory_covers_routes_and_aliases():
    from qa_full_sweep import conversion_inventory

    from kumquat.registry import REGISTRY

    cells = conversion_inventory()
    assert {c["route_id"] for c in cells} == {c.id for c in REGISTRY}
    assert len(cells) == sum(len(c.inputs) for c in REGISTRY)
    assert any(c["source_key"] == ".markdown" for c in cells)
    assert any(c["source_key"] == ".tar.gz" for c in cells)


def test_full_sweep_rejects_empty_office_package(tmp_path):
    import zipfile

    from qa_full_sweep import validate_output

    out = tmp_path / "empty.docx"
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", "<document><body/></document>")
    with pytest.raises(AssertionError):
        validate_output(out, category="documents", source_key=".txt")


def test_full_sweep_rejects_numeric_subtitle_loss(tmp_path):
    from qa_full_sweep import validate_output

    out = tmp_path / "wrong.vtt"
    out.write_text("WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nKumquat\n")
    with pytest.raises(AssertionError):
        validate_output(out, category="subtitles", source_key=".srt")


def test_markdown_alias_converts_real_content(tmp_path):
    from kumquat.apply import apply_conversion

    src = tmp_path / "ملاحظة with spaces.markdown"
    src.write_text("# Kumquat\n\nReal body 42.\n")
    out = apply_conversion("doc-md-to-txt", src, tmp_path / "out")
    assert "Kumquat" in out.read_text()


@pytest.mark.parametrize(
    "family",
    [
        "images",
        "audio",
        "pdf",
        "documents",
        "spreadsheets",
        "subtitles",
        "archives",
        "slides",
    ],
)
def test_conversion_preserves_existing_output(tmp_path, family):
    import zipfile

    import fitz
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from kumquat.apply import apply_conversion

    if family == "images":
        src = tmp_path / "source.png"
        Image.new("RGB", (96, 64), "blue").save(src)
        route, name = "image-png-to-jpg", "source.jpg"
    elif family == "audio":
        src = tmp_path / "source.wav"
        subprocess.run(
            ["ffmpeg", "-v", "error", "-f", "lavfi", "-i", "sine=duration=0.1", str(src)],
            check=True,
            capture_output=True,
        )
        route, name = "audio-wav-to-mp3", "source.mp3"
    elif family == "pdf":
        src = tmp_path / "source.pdf"
        with fitz.open() as doc:
            doc.new_page().insert_text((40, 40), "Kumquat 42")
            doc.save(src)
        route, name = "pdf-to-txt", "source.txt"
    elif family == "documents":
        src = tmp_path / "source.txt"
        src.write_text("Kumquat 42")
        route, name = "doc-txt-to-md", "source.md"
    elif family == "spreadsheets":
        src = tmp_path / "source.csv"
        src.write_text("Kumquat,value\nrow,42\n")
        route, name = "sheet-csv-to-tsv", "source.tsv"
    elif family == "subtitles":
        src = tmp_path / "source.srt"
        src.write_text("1\n00:00:01,000 --> 00:00:02,000\n42\n")
        route, name = "sub-srt-to-vtt", "source.vtt"
    elif family == "archives":
        src = tmp_path / "source.zip"
        with zipfile.ZipFile(src, "w") as archive:
            archive.writestr("hello.txt", "Kumquat")
        route, name = "archive-zip-to-tar", "hello.tar"
    else:
        src = tmp_path / "source.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(6), Inches(1)
        ).text = "Kumquat"
        deck.save(src)
        route, name = "slide-pptx-to-pdf", "source.pdf"
    original = src.read_bytes()
    outdir = tmp_path / "out"
    outdir.mkdir()
    existing = outdir / name
    existing.write_bytes(b"Keep this original output")
    out = apply_conversion(route, src, outdir)
    assert existing.read_bytes() == b"Keep this original output"
    assert out != existing and out.is_file()
    assert src.read_bytes() == original


def test_tar_old_member_timestamp_converts_to_zip(tmp_path):
    import io
    import tarfile
    import zipfile

    from kumquat.apply import apply_conversion

    src = tmp_path / "old.tar"
    with tarfile.open(src, "w") as archive:
        payload = b"Kumquat 42"
        member = tarfile.TarInfo("hello.txt")
        member.mtime = 0
        member.size = len(payload)
        archive.addfile(member, io.BytesIO(payload))
    out = apply_conversion("archive-tar-to-zip", src, tmp_path / "out")
    with zipfile.ZipFile(out) as archive:
        assert archive.read("hello.txt") == b"Kumquat 42"


def test_ods_to_tsv_preserves_existing_csv_without_sidecar(ods, tmp_path):
    from kumquat.converters_docs import convert_sheet

    dest = tmp_path / "out"
    dest.mkdir()
    sentinel = dest / (ods.stem + ".csv")
    sentinel.write_text("keep this output")
    out = convert_sheet(ods, "tsv", dest)
    assert "\t" in out.read_text()
    assert sentinel.read_text() == "keep this output"
    assert sorted(p.suffix for p in dest.iterdir()) == [".csv", ".tsv"]


def test_palette_png_to_jpeg_decodes_pixels(tmp_path):
    from PIL import Image

    from kumquat.converters_images import convert_image

    src = tmp_path / "palette.png"
    im = Image.new("P", (32, 24))
    im.putpalette([200, 80, 20] + [0, 0, 0] * 255)
    im.save(src)
    out = convert_image(src, "jpg", tmp_path / "out")
    with Image.open(out) as decoded:
        assert decoded.size == (32, 24)
        pixel = decoded.getpixel((0, 0))
        assert isinstance(pixel, tuple) and pixel[0] > 190


@pytest.mark.parametrize("dest_ext", ["png", "jpg"])
def test_slide_image_export_preserves_all_slides(tmp_path, dest_ext):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    from kumquat.converters_slides import convert_slide

    src = tmp_path / "شرائح spaced.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    for n in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(2)
        ).text = f"Kumquat slide {n + 1}"
    deck.save(src)
    before = src.read_bytes()
    outdir = tmp_path / "out"
    first = convert_slide(src, dest_ext, outdir)
    images = sorted(outdir.glob(f"*.{dest_ext}"))
    assert len(images) == 2, "multi-slide deck must not silently lose later slides"
    assert first == images[0]
    for path in images:
        with Image.open(path) as image:
            image.load()
            assert image.size == (3000, 1688)
    originals = {p: p.read_bytes() for p in images}
    again = convert_slide(src, dest_ext, outdir)
    assert again not in images
    assert len(list(outdir.glob(f"*.{dest_ext}"))) == 4
    assert all(p.read_bytes() == content for p, content in originals.items())
    assert src.read_bytes() == before


def test_full_sweep_persists_pending_and_completed_cells(tmp_path):
    import json

    from qa_full_sweep import checkpoint

    cells = [
        {"id": "first", "status": "PASS", "reason": "checked"},
        {"id": "second", "status": "UNTESTED", "reason": "Not yet executed"},
    ]
    dest = tmp_path / "progress.json"
    checkpoint(dest, cells, phase="conversions")
    data = json.loads(dest.read_text())
    assert data["complete"] is False
    assert data["conversion_totals"] == {
        "PASS": 1, "FAIL": 0, "UNSUPPORTED": 0, "UNTESTED": 1,
    }
    assert data["conversions"] == cells


@pytest.mark.parametrize("use_runner", [False, True])
def test_full_sweep_timeout_terminates_backend_descendants(tmp_path, use_runner):
    import os
    import sys
    import time

    from qa_full_sweep import exercise, run

    pidfile = tmp_path / "backend.pid"

    def stalled_backend():
        if use_runner:
            run([sys.executable, "-c",
                 "import os,time; from pathlib import Path; "
                 f"Path({str(pidfile)!r}).write_text(str(os.getpid())); time.sleep(60)"])
            return
        proc = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(60)"])
        pidfile.write_text(str(proc.pid))
        proc.wait()

    cell = exercise("edge:timeout", stalled_backend, timeout=0.3)
    assert cell["status"] == "FAIL"
    assert "timeout" in cell["reason"].lower()
    pid = int(pidfile.read_text())
    for _ in range(40):
        stat = __import__("pathlib").Path(f"/proc/{pid}/stat")
        if not stat.exists() or stat.read_text().split()[2] == "Z":
            break
        time.sleep(0.025)
    else:
        os.kill(pid, 9)
        pytest.fail("backend descendant survived timeout")


def test_subtitle_without_cues_is_rejected(tmp_path):
    from kumquat.converters_docs import convert_sub

    src = tmp_path / "invalid.srt"
    src.write_text("Not a subtitle with any timing cues")
    with pytest.raises(ValueError, match="no subtitle cues"):
        convert_sub(src, "vtt", tmp_path / "out")
